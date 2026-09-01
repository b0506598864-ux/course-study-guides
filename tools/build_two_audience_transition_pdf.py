from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_teacher_transition_doc import (
    BLACK,
    BLUE,
    DARK,
    LIGHT_BLUE,
    LIGHT_GOLD,
    MID_GRAY,
    add_bullet,
    add_callout,
    add_heading,
    add_number,
    add_table,
    add_text,
    configure_document,
    set_font,
    set_rtl,
)


ROOT = Path(__file__).resolve().parents[1]
TEMP_DOCX = ROOT / "tmp" / "pdfs" / "transition-two-audiences.docx"


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("מסמך הבהרה לקורס FS AI 45800-9")
    set_font(r, size=11, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph()
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("המעבר משחר לגל\nומה הכיתה צריכה כדי להצליח")
    set_font(r, size=26, bold=True, color=DARK)
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("שני חלקים: הבנה לתלמידים | תמונת מצב למרצה ולרכזת")
    set_font(r, size=13, color=MID_GRAY)
    p.paragraph_format.space_after = Pt(28)

    add_callout(
        doc,
        "משפט אחד שמסכם הכול",
        "הבעיה אינה חוסר רצון ללמוד ואינה חוסר מקצועיות של המרצה. זהו פער מעבר: הכיתה למדה עד עכשיו בדרך הדרגתית ובעברית, וכעת היא נדרשת לעקוב אחרי שפת תעשייה מהירה, אנגלית מקצועית, Git וכלי AI - כאשר רמת הבסיס אינה אחידה.",
        LIGHT_GOLD,
    )
    add_text(
        doc,
        "המסמך מבוסס על תמלול השיעור מ-30.08.2026 ועל חומרי שני חלקי הקורס. הוא נועד להסביר את המצב בלי להאשים תלמידים או מרצים.",
        size=10.5,
        color=MID_GRAY,
        italic=True,
        before=8,
    )


def student_part(doc):
    page_break(doc)
    add_heading(doc, "חלק א: לתלמידים - איך להבין את גל", 1)
    add_callout(doc, "קודם כול", "אתם לא אמורים לדעת כבר את כל המילים שגל משתמש בהן. יש לכם בסיס תכנותי, אבל נוספה שכבה חדשה של כלים, Workflow ואנגלית מקצועית.")

    add_heading(doc, "1. מה השתנה לעומת שחר", 2)
    rows = [
        ("מבנה", "שלב אחר שלב: דרישה, HTML, CSS, JavaScript", "מעבר בין כלי, Git, AI, קוד והסבר תעשייתי"),
        ("שפה", "עברית ופירוק של פקודות", "עברית עם הרבה מונחים באנגלית"),
        ("קצב", "כתיבה הדרגתית שאפשר להעתיק ולבדוק", "הדגמה מהירה, ניסוי, Debug ותיקון בזמן אמת"),
        ("תיעוד", "הערות וקוד ששמרו את רצף השיעור", "ה-Repository הוא מקור האמת, אך לא כל הסבר נאמר גם בכתב"),
        ("AI", "עזר מוגבל למה שנלמד ולדרישות המבחן", "כלי מרכזי: Plan, Agent, Rules ו-Skills"),
        ("הצלחה", "לזהות ולכתוב את ה-Syntax", "להבין דרישה, Context, תכנון ותוצאה"),
    ]
    add_table(doc, ["נושא", "אצל שחר", "אצל גל"], rows, [1700, 3650, 4010])

    add_heading(doc, "2. מה אנחנו יודעים ומה עדיין חדש", 2)
    rows = [
        ("HTML/CSS", "בסיס קיים", "מבנה עמוד, טפסים ועיצוב בסיסי נלמדו"),
        ("JavaScript בסיסי", "בסיס קיים אך דורש חזרה", "משתנים, תנאים, לולאות, פונקציות, מערכים ואובייקטים"),
        ("DOM ו-Forms", "נלמד בסגנון שחר", "צריך כעת להתרגל גם ל-querySelector, createElement ו-addEventListener"),
        ("Local Storage ו-JSON", "מוכר", "התלמידים זיהו stringify, parse ו-Single Source of Truth"),
        ("Git בסיסי", "לא אחיד", "לרוב יש GitHub; לא כולם שולטים ב-clone, pull, branch, merge ו-diff"),
        ("Debugging", "חדש או חלש לרבים", "Breakpoints, Step Into, Scope וקריאת Error"),
        ("Cursor וכלי AI", "חדש", "Ask, Plan, Agent, Context, Rule, Skill ו-Tokens"),
        ("אנגלית מקצועית", "פער משמעותי", "חלק מהתלמידים צריכים קודם לתרגם ורק אחר כך להבין את הפעולה"),
    ]
    add_table(doc, ["תחום", "מצב הכיתה", "המשמעות"], rows, [2100, 2300, 4960])

    add_heading(doc, "3. למה האנגלית והקצב מפריעים", 2)
    add_text(doc, "כאשר גל אומר ברצף Repository, Context, Agent, Breakpoint ו-Relative Path, תלמיד שמכיר אנגלית ממשיך לעקוב אחרי הרעיון. תלמיד שאינו מכיר את המילים צריך לבצע שלוש פעולות במקביל:")
    add_bullet(doc, "לשמוע ולזהות את המילה באנגלית.")
    add_bullet(doc, "לתרגם אותה לעברית.")
    add_bullet(doc, "להבין מה היא עושה בתוך הקוד שרץ כרגע.")
    add_callout(doc, "התוצאה", "אם הקוד כבר התקדם בזמן שהתלמיד עדיין מתרגם את המילה, הוא מאבד את הקשר. זה אינו אומר שהוא אינו מסוגל ללמוד; הוא זקוק לעוד כמה שניות ולמילון קבוע.", LIGHT_GOLD)

    add_heading(doc, "4. פקודה קבועה או שם שהמתכנת בחר?", 2)
    rows = [
        ("const, let, function", "Syntax קבוע של JavaScript", "חייבים לכתוב בצורה שהשפה מכירה"),
        ("querySelector, addEventListener", "מתודות קבועות של ה-DOM", "השם והמבנה קבועים"),
        ("render, init, expenses", "שמות שהמתכנת בוחר", "אפשר לקרוא להם בשם אחר אם הקוד משתמש בו בעקביות"),
        ("Rule, Skill, Agent", "מושגים של כלי AI", "אינם פקודות JavaScript"),
        ("Repository, Commit, Push", "מושגים ופקודות של Git", "אינם JavaScript"),
    ]
    add_table(doc, ["דוגמה", "מה זה", "מה צריך לזכור"], rows, [2600, 3100, 3660])

    add_heading(doc, "5. איך לעבוד בשיעור בלי ללכת לאיבוד", 2)
    for item in (
        "לעצור מיד בשלב שאבד הקשר; לא לחכות לסוף השיעור.",
        "לשאול: האם המילה היא JavaScript, Git, Cursor או שם שהמתכנת בחר?",
        "לרשום בכל שיעור שלוש רשימות: פקודות קבועות, שמות חופשיים, ומילים באנגלית.",
        "אחרי כל הדגמת AI לברר: מה ביקשנו, איזה Context ניתן, מה נוצר ומה נבדק.",
        "לפני שיעורי בית לבצע git pull ולקרוא את README ואת תיקיית השיעור.",
    ):
        add_number(doc, item)
    add_text(doc, "נוסח שאלה שעוזר:", bold=True, color=DARK, before=8, after=2)
    add_text(doc, "עצרתי בשלב ____. הבנתי שאתה מנסה ____. לא הבנתי מה עושה ____. האם זו פקודה קבועה או שם שאתה בחרת? אפשר להראות שוב רק את השלב הזה?", italic=True, after=2)


def coordinator_part(doc):
    add_heading(doc, "חלק ב: למרצה או לרכזת - תמונת מצב ופתרון", 1)
    add_callout(doc, "מטרת החלק", "לא לבקש להאט את כל הקורס או לחזור להתחלה, אלא ליצור גשר קצר בין בסיס הכיתה לבין דרך העבודה החדשה.", LIGHT_GOLD)

    add_heading(doc, "1. תיאור ענייני של המצב", 2)
    add_text(doc, "הכיתה עברה באמצע הקורס בין שתי שיטות הוראה שונות. שחר בנה ידע באמצעות תרגילים קטנים, קוד הדרגתי, הערות בעברית ו-Patterns קבועים. גל מגיע מרקע תעשייתי ומלמד דרך Workflow, Git, Cursor, כלי AI וקוד שנוצר ונבדק בזמן אמת.")
    add_text(doc, "לשתי השיטות יש ערך. הבעיה היא שלא בוצעה שכבת מעבר מספקת בין השיטות, בזמן שרמת האנגלית והניסיון הטכני בכיתה אינם אחידים.")

    add_heading(doc, "2. מה נשמע בפועל בתמלול", 2)
    for item in (
        "תלמיד אמר במפורש שהקצב מהיר ושלא הצליח לעקוב אחרי הפעולות.",
        "תלמידה תיארה את ההסבר כלא מספיק מסודר וביקשה הערות כתובות בתוך חומרי הקורס.",
        "נשמע בלבול בין כלי AI שונים ובין מושגים כגון Rule ו-Skill.",
        "רבים לא הכירו את המונח ECMAScript.",
        "גל זיהה שהכיתה אינה מורגלת ב-Debugging וב-Breakpoints.",
        "נשמע בלבול בין שמות פונקציות שהמתכנת בוחר, כגון render ו-init, לבין פקודות השפה.",
        "ב-Git קיים פער: יש חשבונות ותרגול בסיסי, אך לא כל התלמידים שולטים ב-Workflow מלא.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "3. נקודת המבט של גל", 2)
    add_text(doc, "מהתמלול עולה שגל אינו מתכוון להתעלם מהפערים:")
    add_bullet(doc, "הוא אמר שישלים ויחדד נושאים שלא נלמדו מספיק, אך אינו מתכוון ללמד את הקורס כולו מחדש.")
    add_bullet(doc, "הוא מעדיף השתתפות פעילה, תרגול, שיתוף מסך ושאלות בזמן אמת.")
    add_bullet(doc, "כאשר נאמר לו שהקצב מהיר, הוא עצר וחזר להסבר דרך HTML ו-script שכבר היו מוכרים.")
    add_bullet(doc, "כאשר התבקש לתעד יותר, הוא הסכים לכתוב כשצריך, אף שבתחילה חלק מהדברים נראו לו מובנים מאליהם.")
    add_bullet(doc, "הוא רואה ב-Repository את מקור האמת לכל חומר, הוראה ושיעורי בית.")
    add_callout(doc, "הפער המרכזי", "לגל, כמפתח מנוסה, מונחים רבים הם שפת עבודה טבעית. לתלמיד מתחיל, אותה שפה היא חומר לימוד בפני עצמו. לכן הסבר שנראה למרצה קצר וברור יכול להישמע לתלמיד כרצף של מילים לא מוכרות.")

    add_heading(doc, "4. השונות בין התלמידים", 2)
    rows = [
        ("תלמיד עם אנגלית, Git וכלי AI", "מסוגל לעקוב אחרי הקצב ולהפיק ערך מהדגמות תעשייתיות"),
        ("תלמיד עם בסיס קוד אך אנגלית חלשה", "מבין את הקוד לאחר תרגום, אך מאבד את הרצף בזמן אמת"),
        ("תלמיד מתחיל מאוד", "זקוק להבחנה מפורשת בין Syntax, מונח מקצועי ושם חופשי"),
        ("תלמיד שחושש לומר שלא הבין", "עלול לשתוק; לכן שתיקה בכיתה אינה הוכחה להבנה"),
        ("תלמיד חזק אך לא מסודר", "מבין את ההדגמה בזמן אמת, אך מתקשה לשחזר אותה ללא README וסיכום"),
    ]
    add_table(doc, ["פרופיל למידה", "ההשפעה של המעבר"], rows, [3300, 6060])

    add_heading(doc, "5. הצעה מעשית לגשר של שלושה שיעורים", 2)
    rows = [
        ("לפני השיעור", "README קצר בעברית: מטרות, חמש מילים באנגלית וקבצים שייפתחו"),
        ("בתחילת השיעור", "חמש דקות: מה מוכר מהחלק הקודם ומה חדש היום"),
        ("במהלך ההסבר", "לסמן בקול: זו פקודת JS / זו פקודת Git / זה שם שאני בחרתי"),
        ("כל 20-25 דקות", "Checkpoint קצר: מה ביצענו, מה התוצאה, ושאלה אחת מהכיתה"),
        ("בקוד", "הערות קצרות ליד נקודות מפתח, לא תמלול מלא של כל משפט"),
        ("בסיום", "Commit מסודר + סיכום של עד עשר שורות + שיעורי בית מפורשים"),
        ("לכיתה", "מילון משותף המתעדכן בכל שיעור ומזהי נושאים לשאלות ממוקדות"),
    ]
    add_table(doc, ["מתי", "הפעולה המוצעת"], rows, [2200, 7160])

    add_heading(doc, "6. מה הבקשה מהמרצה ומהרכזת", 2)
    add_text(doc, "הבקשה אינה להוריד את הרמה המקצועית. להפך: לשמור על התוכן התעשייתי, אך להוסיף שכבת תיווך קצרה שתאפשר גם לתלמידים מתחילים להיכנס אליו.")
    add_bullet(doc, "לשמור על Git, Cursor ו-AI כחלק מרכזי מהקורס.")
    add_bullet(doc, "להוסיף מסגרת קבועה בעברית בתחילת ובסיום כל שיעור.")
    add_bullet(doc, "לאמת הבנה באמצעות Checkpoints ולא באמצעות שתיקה.")
    add_bullet(doc, "להבהיר מראש מה חובה לדעת ומה רק הדגמה או העשרה.")
    add_bullet(doc, "לפרסם ב-Repository סיכום ושיעורי בית בנוסח חד-משמעי.")

    add_heading(doc, "7. נוסח קצר שאפשר לומר בפגישה", 2)
    add_callout(
        doc,
        "נוסח מוצע",
        "אנחנו מעריכים את הניסיון המקצועי ואת הכלים החדשים שנכנסו לקורס. הקושי שלנו אינו עצם החומר, אלא המעבר המהיר בין שיטות: עד עכשיו למדנו בשלבים ובעברית, וכעת נכנסו יחד אנגלית מקצועית, Git, Cursor ו-AI. רמת האנגלית והניסיון בכיתה אינה אחידה, ולכן חלק מהתלמידים מאבדים את רצף ההסבר גם כשהם מסוגלים להבין את הנושא. אנחנו מבקשים גשר קצר: מטרות ומילון בתחילת שיעור, Checkpoints במהלך השיעור וסיכום כתוב ב-Repository בסוף. כך נוכל להישאר ברמה המקצועית ולהביא את כל הכיתה איתנו.",
        LIGHT_GOLD,
    )

    add_heading(doc, "מקורות מרכזיים בתמלול", 2)
    for item in (
        "Part 2, 00:09:26-00:11:23 - פערי ידע, השתתפות ושאלות נקודתיות.",
        "Part 2, 00:15:34-00:18:05 - Git Core והמעבר לעבודה דרך Repository.",
        "Part 2, 00:22:52-00:26:35 - Cursor והבלבול בין כלי AI.",
        "Part 2, 02:24:13-02:29:55 - Debugging, הערת הקצב והחזרה להסבר בסיסי.",
        "Part 2, 02:40:03-02:43:03 - בקשת התלמידים לסדר ולהערות כתובות.",
        "Part 2, 02:58:32-03:02:34 - ECMAScript ו-Modern JavaScript.",
        "Part 2, 03:56:28-03:57:14 - render/init כשמות חופשיים וחומר מוכר ב-Local Storage.",
        "Part 1, 00:14:10-00:29:30 - Rules, Skills, Workflow והמעבר מ-Drive ל-Git.",
    ):
        add_bullet(doc, item)
    add_text(doc, "עודכן: 01.09.2026. המסמך אינו כולל שמות תלמידים ואינו מייחס כוונות שלא נאמרו בתמלול.", size=9.5, color=MID_GRAY, italic=True, before=12)


def build():
    doc = Document()
    configure_document(doc)
    cover(doc)
    student_part(doc)
    coordinator_part(doc)
    TEMP_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TEMP_DOCX)
    print(TEMP_DOCX)


if __name__ == "__main__":
    build()
