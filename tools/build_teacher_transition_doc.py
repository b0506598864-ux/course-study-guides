from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "guides" / "share" / "מדריך-מעבר-למורה-החדש.docx"

BLUE = "245A78"
DARK = "17324D"
LIGHT_BLUE = "E8F1F6"
LIGHT_GOLD = "FFF4D6"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
BLACK = "111111"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    # Word needs bidiVisual on the table itself; right-aligned cell text alone
    # does not reverse the visual column order for a Hebrew document.
    bidi_visual = tbl_pr.find(qn("w:bidiVisual"))
    if bidi_visual is None:
        bidi_visual = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi_visual)
    bidi_visual.set(qn("w:val"), "1")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side, value in (("top", 100), ("bottom", 100), ("start", 120), ("end", 120)):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_font(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    rtl = run._element.get_or_add_rPr().find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        run._element.get_or_add_rPr().append(rtl)
    rtl.set(qn("w:val"), "1")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def add_text(doc, text, size=11, bold=False, color=BLACK, after=6, before=0, italic=False):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_rtl(p)
    run = p.add_run(text)
    size = {1: 16, 2: 13, 3: 12}[level]
    color = BLUE if level < 3 else DARK
    set_font(run, size=size, bold=True, color=color)
    p.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}[level])
    p.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}[level])
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_rtl(p)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_rtl(p)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_rtl(p)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True, color=DARK)
    r = p.add_run(text)
    set_font(r, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    for idx, value in enumerate(headers):
        set_cell_shading(header.cells[idx], BLUE)
        p = header.cells[idx].paragraphs[0]
        set_rtl(p)
        r = p.add_run(value)
        set_font(r, size=10.5, bold=True, color=WHITE)
    set_repeat_table_header(header)
    for row_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for col_idx, value in enumerate(row_data):
            if row_idx % 2 == 1:
                set_cell_shading(cells[col_idx], LIGHT_GRAY)
            p = cells[col_idx].paragraphs[0]
            set_rtl(p)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_font(r, size=9.5)
    set_table_geometry(table, widths)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("עמוד ")
    set_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def configure_document(doc):
    section = doc.sections[0]
    sect_pr = section._sectPr
    bidi = sect_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        sect_pr.append(bidi)
    bidi.set(qn("w:val"), "1")
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in (
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK, 10, 5),
    ):
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header
    hp = header.paragraphs[0]
    set_rtl(hp)
    hr = hp.add_run("מדריך מעבר למורה החדש | FS AI 45800-9")
    set_font(hr, size=9, bold=True, color=MID_GRAY)
    add_page_number(section.footer.paragraphs[0])


def build_document():
    doc = Document()
    configure_document(doc)

    # Editorial-cover opening.
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("מדריך קבוצתי")
    set_font(r, size=11, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph()
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("איך ללמוד עם המורה החדש\nבלי ללכת לאיבוד")
    set_font(r, size=27, bold=True, color=DARK)
    p.paragraph_format.space_after = Pt(12)

    p = doc.add_paragraph()
    set_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("מפת מעבר בין שיטת שחר לשיטת גל, המבוססת על התמלולים וקוד השיעור")
    set_font(r, size=13, color=MID_GRAY)
    p.paragraph_format.space_after = Pt(26)

    add_callout(
        doc,
        "המסר המרכזי",
        "הכיתה לא התחילה מחדש, והיא גם לא אמורה כבר להיות צוות פיתוח מקצועי. יש בסיס טוב, אבל רמת הידע אינה אחידה. עכשיו נוספה שכבה חדשה: Git שוטף, כלי AI, אנגלית מקצועית וקריאת קוד שנוצר במהירות.",
        LIGHT_GOLD,
    )

    add_text(doc, "המסמך נועד לכל תלמידי הכיתה. הוא אינו מדרג את המרצים ואינו מחפש אשמים; הוא מסביר מה השתנה ואיך לסגור את הפער בצורה מעשית.", size=10.5, color=MID_GRAY, italic=True, after=18)

    add_heading(doc, "1. מה באמת שמענו בכיתה", 1)
    add_text(doc, "בתמלול נשמעו כמה סימנים ברורים לכך שהמעבר אינו פשוט:")
    add_bullet(doc, "תלמיד עצר את גל ואמר שההסבר מהיר מדי ושלא הצליח לעקוב אחרי הפעולות.")
    add_bullet(doc, "תלמידה תיארה את רצף ההסברים כלא מספיק מסודר וביקשה הערות כתובות, כפי שהיו אצל שחר.")
    add_bullet(doc, "נשאלו שאלות בסיסיות על ההבדלים בין Cursor, Claude, ChatGPT וכלים נוספים.")
    add_bullet(doc, "כאשר גל שאל מי מכיר ECMAScript, התשובה הייתה שרבים אינם מכירים את המונח.")
    add_bullet(doc, "גל זיהה שהכיתה עדיין אינה רגילה ל-Debugging ול-Breakpoints.")
    add_bullet(doc, "כאשר הופיעו שמות כמו render ו-init, תלמיד חשב שאולי אלה פקודות שלא נלמדו; גל הבהיר שאלו שמות שהמתכנת בוחר.")

    add_callout(doc, "המסקנה", "הקושי אינו חוסר יכולת. הוא נובע מפער בשפה, בקצב ובצורת הארגון: גל מדבר לעיתים בשפה של מפתחים עובדים, בזמן שחלק מהכיתה עדיין בונה מילון מושגים בסיסי.")

    add_heading(doc, "2. מה הכיתה כבר יודעת", 1)
    add_text(doc, "לפי תגובות התלמידים והחומר הקודם, יש לכיתה בסיס ממשי בנושאים הבאים:")
    for item in (
        "מבנה HTML וחיבור קובץ JavaScript באמצעות תגית script.",
        "משתנים, תנאים, לולאות, פונקציות, מערכים ואובייקטים ברמה התחלתית.",
        "DOM, טפסים, Events ו-Validation לפי התרגילים של שחר.",
        "Local Storage כמקום שמירת מידע בדפדפן, כולל JSON.stringify ו-JSON.parse.",
        "setTimeout, פעולות אסינכרוניות ועקרונות Promise/async-await שנלמדו בחלק הקודם.",
        "חשבון GitHub ופעולות Git בסיסיות כגון add, commit ו-push; עם זאת, לא כולם תרגלו באותה מידה.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "3. מה עדיין אינו יציב או חדש", 1)
    add_text(doc, "אין פירוש הדבר שאף אחד אינו יודע את הנושאים הבאים. פירוש הדבר שהכיתה כקבוצה עדיין אינה שולטת בהם באופן אחיד:")
    for item in (
        "Workflow מלא של Git: clone, pull, diff, branch, merge ופתרון Conflicts.",
        "Debugging מסודר: Breakpoint, Step Into, Step Over, Scope וקריאת Stack/Error.",
        "עבודה עם Cursor: Ask, Plan, Agent, Context, Rules ו-Skills.",
        "בחירת מודל AI והבנת Tokens, Context ומגבלות שימוש.",
        "מונחי תעשייה באנגלית והבחנה בין מונח קבוע לבין שם שהמתכנת המציא.",
        "קריאת קוד מודרני שנוצר על ידי AI: querySelector, createElement, optional chaining, modules ו-DocumentFragment.",
        "Relative Path לעומת Absolute Path וההשפעה שלהם על הרצת קבצים במחשב אחר.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "4. איך גל מלמד ואיך הוא מגיב", 1)
    rows = [
        ("הקצב", "מהיר ואסוציאטיבי יותר; הוא עובר בין כלי, קוד, Git והסבר תעשייתי."),
        ("הדגמה", "הרבה ניסוי חי: הוא מנסה, נתקל בתקלה, מדבג ומתקן מול הכיתה."),
        ("שאלות", "הוא מבקש לעצור מיד ולשאול שאלה נקודתית. שתיקה עלולה להתפרש אצלו כהבנה."),
        ("פערי ידע", "הוא אמר שיחזור על נושא חסר, אבל אינו מתכוון ללמד את כל הקורס מחדש."),
        ("תגובה לקושי", "כשנאמר לו שהקצב מהיר, הוא עצר וחזר להסבר דרך HTML ו-script tag שכבר מוכרים."),
        ("סדר כתוב", "כשנתבקש לכתוב יותר הערות הוא בתחילה הסביר שהדברים נראים לו פשוטים, אך לבסוף הסכים לתעד כשצריך."),
        ("ציפייה מהתלמיד", "להשתתף, לשתף מסך, לעבוד בפועל ולהגיד במפורש מה לא מובן."),
        ("מקור אמת", "ה-Repository הוא המקום לחומר, להוראות ולשיעורי הבית; לא Drive ולא קוד שנשלח במייל."),
    ]
    add_table(doc, ["נקודה", "מה רואים בתמלול"], rows, [2100, 7260])

    add_callout(doc, "איזון חשוב", "גל אינו מתעלם מהקושי: הוא שאל מה חסר, הבטיח לחדד פערים ועצר כאשר העירו על הקצב. עם זאת, הוא מצפה לשאלה מדויקת בזמן אמת ולא לבקשה לחזור על שיעור שלם בסופו.", LIGHT_GOLD)

    add_heading(doc, "5. ההבדל בין שחר לגל", 1)
    rows = [
        ("מבנה השיעור", "שחר: Syntax ותרגיל שלב אחר שלב", "גל: כלי, Workflow והדגמה תעשייתית"),
        ("שפת ההסבר", "עברית עם פירוק של פקודות", "עברית המשולבת בהרבה מונחים באנגלית"),
        ("כתיבת קוד", "התלמיד כותב לפי Pattern שנלמד", "AI יכול לכתוב; התלמיד צריך לתכנן, לקרוא ולבדוק"),
        ("DOM", "במבחן הועדפו getElementById ו-innerHTML", "ב-Rule החדש מועדפים querySelector ו-createElement"),
        ("Events", "onsubmit בתוך HTML היה חלק מהמבנה", "אין inline events; מחברים addEventListener ב-JS"),
        ("חומר הקורס", "תיקיות, תמלול ו-Drive", "Repository שמתעדכן ב-Git"),
        ("הצלחה בשיעור", "לזכור ולבצע את הפקודות", "להגדיר דרישה, לבחור Context, לבדוק Plan ולהבין את התוצאה"),
    ]
    add_table(doc, ["נושא", "שחר", "גל"], rows, [1700, 3650, 4010])

    add_heading(doc, "6. ההבחנה החשובה: פקודה קבועה או שם שהמתכנת בחר?", 1)
    add_text(doc, "זו אחת מנקודות הבלבול החשובות ביותר. לא כל מילה באנגלית בקוד היא פקודת JavaScript.")
    rows = [
        ("const", "JavaScript", "קבוע", "מכריז על משתנה שלא מבצעים לו הצבה מחדש"),
        ("let", "JavaScript", "קבוע", "מכריז על משתנה שניתן להציב בו ערך חדש"),
        ("function", "JavaScript", "קבוע", "מגדיר פונקציה"),
        ("querySelector", "DOM בדפדפן", "קבוע", "מחפש אלמנט לפי selector"),
        ("addEventListener", "DOM בדפדפן", "קבוע", "מחבר פונקציה לאירוע"),
        ("render", "שם נפוץ", "המתכנת בוחר", "בדרך כלל שם לפונקציה שמציגה מידע"),
        ("init", "שם נפוץ", "המתכנת בוחר", "בדרך כלל שם לפונקציית התחלה"),
        ("expenses", "שם משתנה", "המתכנת בוחר", "מערך או מידע על הוצאות, לפי הקוד"),
        ("Rule", "מושג של Cursor", "מונח כלי", "הוראה קבועה שמנחה את ה-AI"),
        ("Skill", "מושג של כלי AI", "מונח כלי", "Workflow חוזר לביצוע משימה"),
        ("Repository", "Git", "מונח מקצועי", "תיקיית פרויקט עם היסטוריית Git"),
    ]
    add_table(doc, ["מילה", "שייכת ל-", "קבועה או נבחרת?", "פירוש פשוט"], rows, [1800, 1800, 2000, 3760])

    add_heading(doc, "7. מילון הישרדות קצר באנגלית", 1)
    rows = [
        ("Repository / Repo", "מאגר הפרויקט והיסטוריית השינויים"),
        ("Clone", "להוריד עותק של Repository למחשב"),
        ("Commit", "נקודת שמירה מתועדת ב-Git"),
        ("Push", "להעלות Commits למאגר המרוחק"),
        ("Pull", "להוריד עדכונים מהמאגר המרוחק"),
        ("Branch", "ענף עבודה נפרד"),
        ("Merge", "לחבר שינויים מענף אחד לאחר"),
        ("Diff", "ההבדל בין שתי גרסאות"),
        ("Prompt", "הבקשה או ההוראות שאנו נותנים ל-AI"),
        ("Context", "המידע והקבצים שה-AI רואה כרגע"),
        ("Plan", "תוכנית הפעולות לפני שינוי הקוד"),
        ("Agent", "AI שמסוגל לקרוא קבצים ולבצע פעולות"),
        ("Debugging", "איתור והבנת תקלה בקוד"),
        ("Breakpoint", "נקודה שעוצרת את הריצה כדי לבדוק מצב"),
        ("Relative Path", "נתיב ביחס למיקום הקובץ הנוכחי"),
        ("Absolute Path", "נתיב מלא שמתחיל ממיקום קבוע במחשב"),
        ("Inline", "קוד שנכתב בתוך HTML במקום בקובץ נפרד"),
        ("Latest standard", "הגרסה והסגנון העדכניים של השפה"),
    ]
    add_table(doc, ["המונח", "פירוש פשוט בעברית"], rows, [3100, 6260])

    add_heading(doc, "8. איך לשאול את גל שאלה שתעבוד", 1)
    add_text(doc, "במקום לומר רק “לא הבנתי”, משתמשים בתבנית הזאת:")
    add_callout(doc, "תבנית", "עצרתי בשלב ____. הבנתי שאתה מנסה ____. לא הבנתי מה עושה ____. האם זו פקודה קבועה של JavaScript, מושג של Cursor, או שם שאתה בחרת? אפשר להראות שוב רק את השלב הזה?", LIGHT_GOLD)
    add_text(doc, "דוגמה:", bold=True, after=2)
    add_text(doc, "עצרתי כשיצרת את init. הבנתי שהיא מפעילה את תחילת התוכנית, אבל לא הבנתי אם init היא פקודת JavaScript או שם שאתה בחרת. אפשר להראות מאיפה קוראים לה?", italic=True)

    add_heading(doc, "9. שיטת עבודה מומלצת לכל תלמיד", 1)
    for item in (
        "לפני השיעור: לבצע git pull ולפתוח את תיקיית השיעור החדשה.",
        "במהלך השיעור: לרשום שלוש רשימות — פקודות קבועות, שמות שהמתכנת בחר, ומילים באנגלית שעדיין לא ברורות.",
        "כאשר הקצב מהיר: לעצור מיד ולציין את השלב המדויק שבו אבד הקשר.",
        "לאחר הדגמת AI: לשאול מה ה-Prompt, מה ה-Context, מה נוצר ומה נבדק.",
        "לאחר השיעור: לקרוא את README ואת הקבצים החדשים ב-Repository לפני שמתחילים שיעורי בית.",
        "לא להעתיק קוד AI בלי לדעת להסביר את המבנה, הנתונים, האירועים והפונקציות המרכזיות.",
        "כאשר כלל של גל סותר כלל של שחר: לבדוק לאיזו משימה הכלל שייך ולא לערבב ביניהם.",
    ):
        add_number(doc, item)

    add_heading(doc, "10. מה האחריות המשותפת של הכיתה", 1)
    add_bullet(doc, "לומר בכנות כאשר נושא לא נלמד או לא ברור; לא להסתמך על כך שמישהו אחר יעצור.")
    add_bullet(doc, "לבקש מהמרצה שם ברור לכל הדגמה ותוצאה כתובה ב-Repository.")
    add_bullet(doc, "לשתף סיכום ומילון, אך לא להפיץ פרטים אישיים או להפוך את המסמך לביקורת על אדם.")
    add_bullet(doc, "לעזור זה לזה להבחין בין אנגלית מקצועית לבין Syntax אמיתי של השפה.")
    add_bullet(doc, "לרכז שאלות לפני השיעור הבא לפי קובץ, שורה ומזהה — לא כשאלה כללית על כל החומר.")

    add_callout(doc, "השורה התחתונה", "שחר בנה את יסודות השפה; גל מנסה לחבר אותם לדרך עבודה מהתעשייה. המעבר יצליח אם נשמור את ההסברים המסודרים של החלק הראשון, ובמקביל נלמד לעבוד עם Git, אנגלית מקצועית ו-AI בלי לוותר על הבנת הקוד.", LIGHT_GOLD)

    add_heading(doc, "מקורות והפניות בתמלול", 1)
    sources = [
        "Part 2, 00:09:26-00:11:23 — התחייבות להשלים פערים וציפייה לשאלות נקודתיות.",
        "Part 2, 00:15:34-00:18:05 — Git Core, פקודות Git והמעבר לעבודה דרך Repository.",
        "Part 2, 00:22:52-00:26:35 — בחירת Cursor והבלבול בין כלי AI שונים.",
        "Part 2, 02:24:13-02:27:33 — פער ב-Debugging והערה שהקצב מהיר.",
        "Part 2, 02:27:33-02:29:55 — עצירה וחזרה להסבר דרך HTML, script ונתיבים.",
        "Part 2, 02:40:03-02:43:03 — בקשת התלמידים לסדר ולהערות כתובות ותגובת המרצה.",
        "Part 2, 02:58:32-03:02:34 — ECMAScript ו-Modern JavaScript.",
        "Part 2, 03:56:28-03:57:14 — render ו-init כשמות שהמתכנת בוחר; Local Storage כחומר מוכר.",
        "Part 1, 00:14:10-00:27:19 — Rules, Skills ו-Workflow של push.",
        "Part 1, 00:28:54-00:29:30 — חומר ושיעורי בית ב-Repository ולא ב-Drive.",
    ]
    for item in sources:
        add_bullet(doc, item)

    add_text(doc, "הוכן על בסיס תמלול השיעור מ-30.08.2026 וקובצי ה-Repository של הקורס. עודכן: 01.09.2026.", size=9.5, color=MID_GRAY, italic=True, before=14)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
